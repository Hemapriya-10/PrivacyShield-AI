"""
Document Extraction Module for PrivacyShield AI.

Provides functions to safely extract plain text from TXT, PDF, and DOCX file byte streams
without persisting documents to local storage.
"""

import io
from typing import Tuple


def validate_file_type(filename: str) -> str:
    """
    Validates file extension and returns the lowercase extension without leading dot.
    
    Raises:
        ValueError: If file extension is unsupported or filename is empty.
    """
    if not filename or '.' not in filename:
        raise ValueError("Invalid filename. File must have a valid extension (.txt, .pdf, .docx).")
    
    ext = filename.rsplit('.', 1)[-1].lower()
    allowed_extensions = {'txt', 'pdf', 'docx'}
    if ext not in allowed_extensions:
        raise ValueError(
            f"Unsupported file format '.{ext}'. Supported formats are: .txt, .pdf, .docx"
        )
    return ext


def extract_from_txt(file_bytes: bytes) -> str:
    """Extracts text from TXT byte stream with UTF-8 and Latin-1 fallbacks."""
    if not file_bytes:
        raise ValueError("The uploaded TXT document is empty.")
    
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = file_bytes.decode("latin-1")
        except Exception as e:
            raise ValueError(f"Failed to decode TXT file content: {str(e)}")
            
    if not text.strip():
        raise ValueError("The uploaded TXT document contains no readable text.")
        
    return text.strip()


def extract_from_pdf(file_bytes: bytes) -> str:
    """Extracts text from PDF byte stream using PyMuPDF (fitz)."""
    if not file_bytes:
        raise ValueError("The uploaded PDF document is empty.")
        
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise ImportError("PyMuPDF is required for PDF extraction. Please install fitz / PyMuPDF.")

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise ValueError(f"Failed to open or parse PDF file: Corrupted or invalid format ({str(e)})")

    extracted_pages = []
    try:
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            page_text = page.get_text()
            if page_text:
                extracted_pages.append(page_text.strip())
        doc.close()
    except Exception as e:
        doc.close()
        raise ValueError(f"Error while reading PDF pages: {str(e)}")

    full_text = "\n\n".join(extracted_pages).strip()
    if not full_text:
        raise ValueError("The PDF file was processed successfully, but no readable text was found.")

    return full_text


def extract_from_docx(file_bytes: bytes) -> str:
    """Extracts text from DOCX byte stream using python-docx."""
    if not file_bytes:
        raise ValueError("The uploaded DOCX document is empty.")

    try:
        import docx
    except ImportError:
        raise ImportError("python-docx is required for DOCX extraction. Please install python-docx.")

    try:
        doc_stream = io.BytesIO(file_bytes)
        document = docx.Document(doc_stream)
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX file: Corrupted or invalid format ({str(e)})")

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    
    # Also extract text from tables if present
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                paragraphs.append(row_text)

    full_text = "\n\n".join(paragraphs).strip()
    if not full_text:
        raise ValueError("The DOCX file was processed successfully, but no readable text was found.")

    return full_text


def extract_text(file_bytes: bytes, filename: str) -> str:
    """
    Main entry point for document text extraction.
    
    Args:
        file_bytes (bytes): Raw binary content of the file.
        filename (str): Name of the file (used for extension detection).
        
    Returns:
        str: Extracted plain text string.
        
    Raises:
        ValueError: On unsupported format, empty file, or parsing failure.
    """
    ext = validate_file_type(filename)
    
    if ext == 'txt':
        return extract_from_txt(file_bytes)
    elif ext == 'pdf':
        return extract_from_pdf(file_bytes)
    elif ext == 'docx':
        return extract_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported extension: .{ext}")
