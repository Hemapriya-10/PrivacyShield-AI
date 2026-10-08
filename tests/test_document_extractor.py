"""
Unit tests for Document Extraction Module.
"""

import io
import pytest
from modules.document_extractor import (
    validate_file_type,
    extract_from_txt,
    extract_from_pdf,
    extract_from_docx,
    extract_text,
)


def test_validate_file_type():
    """Test file extension validation."""
    assert validate_file_type("document.txt") == "txt"
    assert validate_file_type("report.PDF") == "pdf"
    assert validate_file_type("file.docx") == "docx"

    with pytest.raises(ValueError, match="Unsupported file format"):
        validate_file_type("script.py")

    with pytest.raises(ValueError, match="Invalid filename"):
        validate_file_type("noextension")


def test_extract_from_txt():
    """Test text extraction from TXT bytes."""
    content = "Hello PrivacyShield AI!\nThis is a sample document."
    txt_bytes = content.encode("utf-8")
    extracted = extract_from_txt(txt_bytes)
    assert extracted == content

    # Empty file test
    with pytest.raises(ValueError, match="empty"):
        extract_from_txt(b"")

    with pytest.raises(ValueError, match="no readable text"):
        extract_from_txt(b"   \n  \t  ")


def test_extract_from_pdf():
    """Test text extraction from in-memory PDF bytes."""
    try:
        import fitz
    except ImportError:
        pytest.skip("PyMuPDF not installed")

    # Create a simple PDF in memory
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "PDF Test Document Content")
    pdf_bytes = doc.tobytes()
    doc.close()

    extracted = extract_from_pdf(pdf_bytes)
    assert "PDF Test Document Content" in extracted

    # Corrupted PDF test
    with pytest.raises(ValueError, match="Failed to open or parse PDF"):
        extract_from_pdf(b"Not a valid PDF binary content")


def test_extract_from_docx():
    """Test text extraction from in-memory DOCX bytes."""
    try:
        import docx
    except ImportError:
        pytest.skip("python-docx not installed")

    # Create a simple DOCX in memory
    doc = docx.Document()
    doc.add_paragraph("DOCX Test Paragraph One")
    doc.add_paragraph("DOCX Test Paragraph Two")
    
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    extracted = extract_from_docx(docx_bytes)
    assert "DOCX Test Paragraph One" in extracted
    assert "DOCX Test Paragraph Two" in extracted

    # Corrupted DOCX test
    with pytest.raises(ValueError, match="Failed to parse DOCX file"):
        extract_from_docx(b"Not a valid docx content")


def test_extract_text_wrapper():
    """Test main extract_text wrapper."""
    txt_bytes = "Test content".encode("utf-8")
    res = extract_text(txt_bytes, "sample.txt")
    assert res == "Test content"
